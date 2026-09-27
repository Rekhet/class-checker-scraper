from __future__ import annotations

import csv
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from scraper import changelog, export


def _cls(name, *, prof="김교수", slots=(("09:00", "10:15"),), cls=("학사", "전선"), quota=30):
    return {"name": name, "professor": prof, "classification": list(cls), "quota": quota,
            "credits": 3, "room": "43-101", "department": "통계학과",
            "slots": [{"day_index": 0, "start_time": a, "end_time": b} for a, b in slots]}


K = ("2026", "T1")


class ChangelogDiffTests(unittest.TestCase):
    def test_new_removed_and_field_changes_are_coded(self) -> None:
        old = {(*K, "A", "001"): _cls("가"), (*K, "B", "001"): _cls("나"),
               (*K, "C", "001"): _cls("다")}
        new = {(*K, "A", "001"): _cls("가", prof="이교수", slots=(("13:00", "14:15"),)),
               (*K, "C", "001"): _cls("다"),
               (*K, "D", "001"): _cls("라", slots=())}
        rows, counts = changelog.diff(old, new)
        codes = sorted((code, key[2]) for code, key, _name, _d in rows)
        self.assertIn(("NEW", "D"), codes)
        self.assertIn(("DEL", "B"), codes)
        self.assertIn(("TCHG", "A"), codes)
        self.assertTrue(any(k == "A" and c not in ("TCHG",) for c, k in codes))   # professor
        self.assertEqual((counts["new"], counts["removed"], counts["changed"]), (1, 1, 1))
        self.assertNotIn("C", {k for _c, k in codes})              # unchanged class

    def test_empty_and_none_are_the_same(self) -> None:
        a = _cls("가")
        b = {**a, "room": None}
        a["room"] = ""
        self.assertEqual(changelog.diff({(*K, "A", "1"): a}, {(*K, "A", "1"): b})[0], [])

    def test_run_log_is_written_with_a_header(self) -> None:
        with tempfile.TemporaryDirectory() as tmp, \
             patch.object(changelog, "LOG_DIR", Path(tmp)):
            rows, counts = changelog.diff({}, {(*K, "A", "001"): _cls("가")})
            path = changelog.write_run_log(7, rows, counts)
            text = path.read_text(encoding="utf-8")
        self.assertIn("summary: new=1 removed=0 changed=0", text)
        self.assertIn("NEW\t2026/T1\tA(001)\t가", text)


class ExportTests(unittest.TestCase):
    ROWS = [{**_cls("통계학"), "year": "2026", "term": "U000200002U000300001",
             "sbjt_cd": "326.211", "lt_no": "001", "applied": 10, "enrolled": 9}]

    def test_csv_opens_in_excel_as_utf8(self) -> None:
        data = export.to_csv(self.ROWS)
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))          # BOM
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        self.assertEqual(len(rows), 2)
        self.assertIn("통계학", rows[1])

    def test_xlsx_is_a_workbook_with_the_same_rows(self) -> None:
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl (export extra) not installed")
        wb = openpyxl.load_workbook(io.BytesIO(export.to_xlsx(self.ROWS)))
        ws = wb.active
        self.assertEqual(ws.max_row, 2)
        self.assertIn("통계학", [c.value for c in ws[2]])


if __name__ == "__main__":
    unittest.main()
